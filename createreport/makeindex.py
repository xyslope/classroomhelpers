from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_TAB_ALIGNMENT, WD_TAB_LEADER
import sys
import os
import shutil
import pandas as pd

def add_toc_entry(doc, keyword, page):
    """
    キーワードとページ番号の間にリーダー（点々）を入れ、ページ番号を右寄せする目次エントリーを追加
    """
    # 段落を作成
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT  # 左揃え

    # タブストップの設定
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)  # 右揃え + 点々リーダー

    # キーワードを追加
    run = paragraph.add_run(keyword)
    run.font.size = Pt(12)

    # タブを挿入
    paragraph.add_run("\t")  # タブ文字でリーダーを挿入

    # ページ番号を追加
    run_page = paragraph.add_run(str(page))
    run_page.font.size = Pt(12)

if __name__ == "__main__":
    # get_options
    args = get_options()
    source_path = args.source_path
    docdir = Path(source_path) / args.doc_dir.lstrip("\\/")
    if not docdir.exists():
        print('ソースディレクトリが存在しません。')
        sys.exit()
 
    # create filelist
    #    filelist = pd.read_csv(Path(source_path) / 'filelist.csv')
    filelist =  str(Path(source_path) / 'filelist2.csv')

# Word 文書を作成
doc = Document()
doc.add_heading("目次", level=1)

# 目次エントリーを追加
entries = [("第一章: はじめに", 1), ("第二章: 理論", 5), ("第三章: 実験", 12), ("第四章: 結果", 20)]
for keyword, page in entries:
    add_toc_entry(doc, keyword, page)

# 保存
doc.save("toc_with_leaders.docx")
